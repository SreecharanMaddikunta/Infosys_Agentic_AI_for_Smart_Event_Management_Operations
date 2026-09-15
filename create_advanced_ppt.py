import fitz
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import glob

# Paths to the 4 milestone PDFs
pdf_files = {
    "M1": r"C:\Users\vinay\Downloads\Infosys_MileStone_1_Sreecharan_Maddikunta.pdf",
    "M2": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_2_.pdf",
    "M3": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_3_final_Document (1).pdf",
    "M4": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_4_Final_Documentation (1).pdf"
}

output_img_dir = "extracted_screenshots"
os.makedirs(output_img_dir, exist_ok=True)

extracted_images = {"M1": [], "M2": [], "M3": [], "M4": []}

for milestone, path in pdf_files.items():
    if not os.path.exists(path):
        print(f"Warning: File {path} not found.")
        continue
    try:
        doc = fitz.open(path)
        img_count = 0
        for page_num in range(len(doc)):
            page = doc[page_num]
            image_list = page.get_images(full=True)
            for img_index, img in enumerate(image_list):
                xref = img[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]
                image_ext = base_image["ext"]
                
                # Filter out small images (logos, icons)
                if len(image_bytes) < 15000:
                    continue
                    
                image_path = os.path.join(output_img_dir, f"{milestone}_page{page_num}_{img_index}.{image_ext}")
                with open(image_path, "wb") as f:
                    f.write(image_bytes)
                extracted_images[milestone].append(image_path)
                img_count += 1
        print(f"Extracted {img_count} large images from {milestone}.")
    except Exception as e:
        print(f"Error extracting from {milestone}: {e}")


# Initialize PPT
prs = Presentation()

def apply_blue_theme(shape):
    if not shape.has_text_frame: return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 85, 165) # Infosys Blue

def add_title_slide(title_text, subtitle_text):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title_text
    slide.placeholders[1].text = subtitle_text
    apply_blue_theme(slide.shapes.title)

def add_slide_with_image(title_text, bullets, img_path=None, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title only
    
    # Title
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_shape.left = Inches(0.5)
    title_shape.top = Inches(0.3)
    title_shape.width = Inches(9)
    title_shape.height = Inches(1)
    apply_blue_theme(title_shape)
    
    # Text Box (Left)
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(4.5)
    height = Inches(5.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(12)
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(60, 60, 60)
        
    # Image (Right)
    if img_path and os.path.exists(img_path):
        try:
            img_left = Inches(5.2)
            img_top = Inches(1.5)
            max_width = Inches(4.3)
            slide.shapes.add_picture(img_path, img_left, img_top, width=max_width)
        except Exception as e:
            print(f"Failed to add image {img_path}: {e}")

    # Notes
    if notes:
        slide.notes_slide.notes_text_frame.text = notes

def get_img(milestone, index=-1):
    imgs = extracted_images.get(milestone, [])
    if not imgs:
        return None
    # Wrap index
    index = index % len(imgs)
    return imgs[index]

# Build the presentation
add_title_slide("Agentic AI for Smart Event Management Operations", 
                "Infosys Springboard Virtual Summer Internship 7.0\nPresented by: Sreecharan Maddikunta")

add_slide_with_image("The Paradigm Shift in Event Operations",
    ["Moving from reactive to proactive, correlated intelligence.", "Replacing manual scheduling with goal-oriented autonomous agents.", "Bridging the gap between planning, live execution, and sponsor ROI."],
    None, "Traditional event management relies on spreadsheets and human intuition...")

add_slide_with_image("Milestone 1: Intelligent Registration",
    ["Zero-shot LLM Spam & Identity Validation.", "Dynamic Waitlist Engine wrapped in Prisma transactions.", "Automated HTML Email Ticketing with dynamic QR codes."],
    get_img("M1", -3), "Milestone 1 focused on the attendee experience...")

add_slide_with_image("Live Analytics & Real-Time VIP Alerts",
    ["Server-side demographic aggregations via Recharts.", "HTML5 QR scanning for physical venue check-in.", "Real-time VIP arrival alerts pushed via Socket.io."],
    get_img("M1", -1), "On the operations side, the Admin Dashboard provides real-time demographic analytics...")

add_slide_with_image("Milestone 2: Agentic Venue & Speaker Operations",
    ["Transitioning from validation to goal-oriented planning.", "Combining deterministic constraints with LLM semantic matching.", "Human-in-the-loop: Agents recommend, Admins approve."],
    get_img("M2", -4), "Milestone 2 introduced our first goal-oriented agents...")

add_slide_with_image("Agentic Optimization & Scheduling",
    ["Venue Optimization: Weights capacity, equipment, and accessibility.", "Speaker Scheduling: Matches bio/expertise against session topics.", "Strict schema validation ensures no hallucinated assignments reach the database."],
    get_img("M2", -1), "Here we see the agents in action...")

add_slide_with_image("Milestone 3: AI Incident Triage Agent",
    ["Free-text natural language reporting by students.", "AI automatically classifies Severity (Low to Critical) and Category.", "Smart routing to specific operational teams (Medical, Security, Tech)."],
    get_img("M3", -4), "Milestone 3 shifted focus to live, safety-critical execution...")

add_slide_with_image("Live Sponsor Engagement & ROI Tracking",
    ["Brand-new, isolated React application for Sponsors.", "Agent computes real-time ROI based on live booth traffic.", "Generates natural-language performance summaries and alerts."],
    get_img("M3", -1), "Sponsors need to know their investment is paying off...")

add_slide_with_image("Milestone 4: Proactive Event Intelligence",
    ["Continuous monitoring of live operational telemetry.", "Detects anomalies (e.g., registration queue overflow, hall capacity limits).", "Generates EventInsight and ActionableRecommendation records."],
    get_img("M4", -2), "Milestone 4 elevates the system to an enterprise level...")

add_slide_with_image("Executive Command Center",
    ["Side-by-side view of anomalies and AI-proposed solutions.", "One-click Accept / Dismiss actions for rapid triage.", "Full audit logging of administrative decisions."],
    get_img("M4", -1), "All this intelligence culminates in the Executive Command Center...")

add_slide_with_image("AI Agents Fleet Visualization",
    ["Complete transparency into agent capabilities.", "Displays JSON input/output schemas and reasoning logic.", "Demystifies the 'black box' of LLM operations."],
    get_img("M4", -3), "To ensure transparency and trust, I built the AI Agents Visualization Dashboard...")

add_slide_with_image("Conclusion",
    ["Delivered a cohesive, 5-service enterprise platform.", "Successfully merged generative AI with deterministic software engineering.", "Achieved true Agentic operational intelligence."],
    None, "In conclusion, this internship culminated in a highly complex, 5-service enterprise platform.")

output_path = r"C:\Users\vinay\Desktop\InfosysProject1\Agentic_AI_Event_Management_Advanced.pptx"
prs.save(output_path)
print(f"Advanced PPTX successfully saved to {output_path}")

# Cleanup
for f in glob.glob(os.path.join(output_img_dir, "*")):
    os.remove(f)
os.rmdir(output_img_dir)
print("Temporary images cleaned up.")
