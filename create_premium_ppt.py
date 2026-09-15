import os
import glob
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

pdf_files = {
    "M1": r"C:\Users\vinay\Downloads\Infosys_MileStone_1_Sreecharan_Maddikunta.pdf",
    "M2": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_2_.pdf",
    "M3": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_3_final_Document (1).pdf",
    "M4": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_4_Final_Documentation (1).pdf"
}

output_img_dir = "ext_imgs_premium"
os.makedirs(output_img_dir, exist_ok=True)
extracted_images = {"M1": [], "M2": [], "M3": [], "M4": []}

for milestone, path in pdf_files.items():
    if not os.path.exists(path):
        continue
    try:
        doc = fitz.open(path)
        temp_imgs = []
        for page_num in range(len(doc)):
            image_list = doc[page_num].get_images(full=True)
            for img_index, img in enumerate(image_list):
                base_image = doc.extract_image(img[0])
                image_bytes = base_image["image"]
                if len(image_bytes) < 40000: # High threshold to skip small diagrams/logos/code snippets
                    continue
                image_path = os.path.join(output_img_dir, f"{milestone}_p{page_num:03d}_{img_index}.{base_image['ext']}")
                with open(image_path, "wb") as f:
                    f.write(image_bytes)
                temp_imgs.append(image_path)
        # We only want the LAST 6 images because those are the UI dashboards
        extracted_images[milestone] = temp_imgs[-6:] if len(temp_imgs) >= 6 else temp_imgs
    except Exception as e:
        print(f"Error extracting {milestone}: {e}")

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# Brand Colors
NAVY = RGBColor(15, 23, 42)      # Midnight Blue
SLATE = RGBColor(100, 116, 139)  # Slate Gray
ACCENT = RGBColor(0, 102, 204)   # Infosys Blue
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(248, 250, 252)

def add_footer(slide):
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(12), Inches(0.3))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Infosys Springboard 7.0 | Agentic AI for Smart Event Management Operations"
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE
    p.font.name = "Segoe UI"
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.08)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = ACCENT
    slide.shapes[-1].line.color.rgb = ACCENT

def add_title_slide(title="Agentic AI for Smart Event\nManagement Operations", subtitle="Infosys Springboard Virtual Summer Internship 7.0\nPresented by: Sreecharan Maddikunta"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(22)
    p2.font.color.rgb = SLATE
    p2.font.name = "Segoe UI"
    p2.alignment = PP_ALIGN.CENTER

def add_text_slide(title, subtitle, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_footer(slide)
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.color.rgb = NAVY
    p.font.bold = True
    p.font.name = "Segoe UI"
    
    if subtitle:
        p2 = txBox.text_frame.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = ACCENT
        p2.font.name = "Segoe UI"

    left = Inches(0.5)
    top = Inches(1.8)
    width = Inches(12)
    height = Inches(5.0)
    txBox2 = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox2.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + b
        p.level = 0
        p.space_after = Pt(16)
        p.font.size = Pt(20)
        p.font.color.rgb = NAVY
        p.font.name = "Segoe UI"

def add_flowchart_slide(title, subtitle, steps):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()
    add_footer(slide)
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.color.rgb = NAVY
    p.font.bold = True
    p.font.name = "Segoe UI"
    
    if subtitle:
        p2 = txBox.text_frame.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(20)
        p2.font.color.rgb = ACCENT
        p2.font.name = "Segoe UI"

    box_w = 2.4
    box_h = 1.3
    spacing = 0.6
    start_top = 3.0
    
    total_w = (len(steps) * box_w) + ((len(steps) - 1) * spacing)
    start_left = (13.33 - total_w) / 2.0

    for i, step in enumerate(steps):
        l = start_left + (i * (box_w + spacing))
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(start_top), Inches(box_w), Inches(box_h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = NAVY
        rect.line.color.rgb = ACCENT
        rect.line.width = Pt(3)
        tf = rect.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = step
        p.font.size = Pt(18)
        p.font.color.rgb = WHITE
        p.font.name = "Segoe UI"
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(l + box_w + 0.1), Inches(start_top + 0.45), Inches(spacing - 0.2), Inches(0.4))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT
            arrow.line.fill.background()

def add_ui_showcase(title, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    add_footer(slide)
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "Segoe UI"

    # 3 Image Grid
    img_w = 3.8
    img_h = 4.2
    spacing = 0.3
    start_left = (13.33 - (3 * img_w + 2 * spacing)) / 2.0
    start_top = 1.8
    
    valid_imgs = [img for img in images if img and os.path.exists(img)]
    for i, img in enumerate(valid_imgs[:3]):
        try:
            l = start_left + (i * (img_w + spacing))
            slide.shapes.add_picture(img, Inches(l), Inches(start_top), width=Inches(img_w))
        except Exception as e:
            pass

# --- SLIDE CREATION ---
add_title_slide()

add_text_slide("Project Overview & Vision", "The Paradigm Shift in Event Operations", [
    "Moving from reactive to proactive, correlated intelligence.",
    "Replacing manual scheduling with goal-oriented autonomous agents.",
    "Bridging the gap between planning, live execution, and sponsor ROI.",
    "Eliminating human bottlenecks with secure, human-in-the-loop approvals."
])

add_flowchart_slide("Global System Architecture", "Decoupled Event-Driven Microservices", [
    "React Frontends\n(Student/Admin)", "Node.js / Express\n(Orchestration)", "PostgreSQL\n(Prisma ORM)", "Python / FastAPI\n(Agentic LLMs)"
])

def create_milestone(m_id, m_name, overview, tech, flow_title, flow_steps):
    add_text_slide(f"Milestone {m_id}: {m_name}", "Milestone Overview", overview)
    add_text_slide(f"Milestone {m_id}: Technical Architecture", "Technologies & Engineering Principles", tech)
    add_flowchart_slide(f"Milestone {m_id}: Workflow", flow_title, flow_steps)
    
    imgs = extracted_images.get(f"M{m_id}", [])
    m_imgs = imgs[-6:] if len(imgs) >= 6 else (imgs + [None]*6)[:6]
    add_ui_showcase(f"Milestone {m_id}: Output Showcase 1", m_imgs[0:3])
    add_ui_showcase(f"Milestone {m_id}: Output Showcase 2", m_imgs[3:6])

create_milestone(1, "Intelligent Registration",
    ["Zero-shot LLM Spam & Identity Validation blocks fraudulent registrations.", "Dynamic Waitlist Engine prevents capacity race conditions.", "Automated HTML Email Ticketing dispatches dynamic QR codes.", "Real-Time VIP Alerts trigger sub-second toasts via Socket.io."],
    ["Frontend: React.js, TailwindCSS", "Backend: Node.js, Express, Socket.io", "Database: PostgreSQL, Prisma", "AI: Python, FastAPI, Gemini Flash"],
    "Registration & Validation Pipeline",
    ["Student Form\nSubmission", "AI Semantic\nValidation", "Waitlist Capacity\nCheck (SQL)", "QR Ticket\nGeneration"]
)

create_milestone(2, "Venue & Speaker Operations",
    ["Transitioning from validation to goal-oriented operational planning.", "Venue Optimization weights capacity, equipment, and accessibility.", "Speaker Scheduling matches bio/expertise against session topics.", "Strict schema validation prevents hallucinated database assignments."],
    ["Agent Architecture: Python / FastAPI", "Orchestration: Node.js Controllers", "Validation: Pydantic Strict Schemas", "Constraints: SQL Deterministic Locks"],
    "Agentic Optimization Workflow",
    ["Context Retrieval\n(Venues)", "Hard Constraint\nFiltering", "LLM Soft\nScoring", "Admin Review &\nPersistence"]
)

create_milestone(3, "Sponsorship & Incident Agents",
    ["Live Incident Triage Agent categorizes free-text emergency reports.", "Smart routing to specific operational teams (Medical, Security).", "Live Sponsor Engagement & ROI Tracking via dedicated portal.", "High-priority operational alerts pushed to Socket.io."],
    ["Real-Time: Socket.io Client/Server", "AI: Incident & Sponsor Agents", "Frontend: Isolated Sponsor React App", "Security: Hardened JWT Expiration"],
    "Live Incident Triage Pipeline",
    ["Vague Incident\nReport", "LLM Severity\nClassification", "Socket.io\n'new_alert'", "Live Kanban\nBoard"]
)

create_milestone(4, "Event Intelligence & Deployment",
    ["Proactive Monitoring detects anomalies (e.g., queue overflow).", "Actionable Recommendations provide one-click 'Accept' solutions.", "Multi-Agent Orchestrator delegates tasks to Venue/Sponsor agents in parallel.", "Strict RBAC closes privilege gaps for enterprise deployment."],
    ["Orchestration: Agent Orchestrator Logic", "Intelligence: Telemetry Monitoring", "Security: Strict RBAC Headers", "Deployment: Unified Concurrent Pipeline"],
    "Multi-Agent Orchestration Workflow",
    ["Complex Domain\nProblem", "Task\nDecomposition", "Parallel Agent\nDelegation", "Aggregated\nAction Plan"]
)

add_text_slide("Results & Evaluation", "Platform Performance Metrics", [
    "Reliability: Zero hallucinated database writes due to strict Pydantic validation.",
    "Performance: Real-time Socket.io alerts achieved sub-second latency.",
    "Concurrency: Database transactions fully resolved waitlist race conditions.",
    "Security: Complete isolation between Sponsor and Admin authorization."
])

add_text_slide("Future Enhancements", "The Roadmap Ahead", [
    "Predictive Attendance Modeling for automated venue scaling.",
    "Mobile Push/SMS notifications for on-duty field staff.",
    "Incident Heatmaps for spatial pattern detection.",
    "Voice-based administration via NLP commands."
])

add_title_slide("Thank You", "Agentic AI for Smart Event Management Operations\nReady for Enterprise Deployment")

output_path = r"C:\Users\vinay\Desktop\InfosysProject1\Agentic_AI_Event_Management_Premium.pptx"
prs.save(output_path)
print(f"Premium PPTX successfully saved to {output_path}")

for f in glob.glob(os.path.join(output_img_dir, "*")):
    try:
        os.remove(f)
    except:
        pass
try:
    os.rmdir(output_img_dir)
except:
    pass
print("Temporary images cleaned up.")
