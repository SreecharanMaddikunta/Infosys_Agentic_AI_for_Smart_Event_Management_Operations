import os
import glob
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --- 1. IMAGE EXTRACTION ---
pdf_files = {
    "M1": r"C:\Users\vinay\Downloads\Infosys_MileStone_1_Sreecharan_Maddikunta.pdf",
    "M2": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_2_.pdf",
    "M3": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_3_final_Document (1).pdf",
    "M4": r"C:\Users\vinay\Downloads\Sreecharan_Maddikunta_Milestone_4_Final_Documentation (1).pdf"
}

output_img_dir = "ext_imgs"
os.makedirs(output_img_dir, exist_ok=True)
extracted_images = {"M1": [], "M2": [], "M3": [], "M4": []}

for milestone, path in pdf_files.items():
    if not os.path.exists(path):
        continue
    try:
        doc = fitz.open(path)
        for page_num in range(len(doc)):
            image_list = doc[page_num].get_images(full=True)
            for img_index, img in enumerate(image_list):
                base_image = doc.extract_image(img[0])
                image_bytes = base_image["image"]
                if len(image_bytes) < 20000: # Filter small logos
                    continue
                image_path = os.path.join(output_img_dir, f"{milestone}_p{page_num}_{img_index}.{base_image['ext']}")
                with open(image_path, "wb") as f:
                    f.write(image_bytes)
                extracted_images[milestone].append(image_path)
    except Exception as e:
        print(f"Error extracting {milestone}: {e}")

def get_img(milestone, index=0):
    imgs = extracted_images.get(milestone, [])
    if not imgs: return None
    return imgs[index % len(imgs)]

# --- 2. PPTX GENERATION & STYLING ---
prs = Presentation()
prs.slide_width = Inches(13.33) # 16:9 widescreen
prs.slide_height = Inches(7.5)

# Brand Colors
NAVY = RGBColor(0, 32, 80)
BLUE = RGBColor(0, 102, 204)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(230, 230, 230)
DARK_GRAY = RGBColor(60, 60, 60)

def add_footer(slide):
    left = Inches(0.5)
    top = Inches(7.1)
    width = Inches(12)
    height = Inches(0.3)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Infosys Springboard 7.0 | Agentic AI for Smart Event Management Operations"
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(150, 150, 150)
    
    # Top border line
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.1)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BLUE
    slide.shapes[-1].line.color.rgb = BLUE

def add_title_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank
    # Dark Navy Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Agentic AI for Smart Event\nManagement Operations"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Infosys Springboard Virtual Summer Internship 7.0\nPresented by: Sreecharan Maddikunta"
    p2.font.size = Pt(24)
    p2.font.color.rgb = GRAY
    p2.alignment = PP_ALIGN.CENTER

def add_section_header(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(24)
        p2.font.color.rgb = GRAY
        p2.alignment = PP_ALIGN.CENTER

def add_content_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_footer(slide)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.color.rgb = NAVY
    p.font.bold = True
    
    # Content
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(12)
    height = Inches(5.5)
    txBox2 = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox2.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.space_after = Pt(14)
        p.font.size = Pt(22)
        p.font.color.rgb = DARK_GRAY

def add_dashboard_showcase(title, descriptions, img_paths):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_footer(slide)
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.color.rgb = NAVY
    p.font.bold = True
    
    # Left text column
    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4), Inches(5.5))
    tf = txBox2.text_frame
    tf.word_wrap = True
    for i, b in enumerate(descriptions):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.space_after = Pt(14)
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY

    # Add images dynamically in a grid on the right
    img_left = 5.0
    img_top = 1.5
    img_max_w = 7.5
    img_max_h = 5.0
    
    valid_imgs = [img for img in img_paths if img and os.path.exists(img)]
    
    if len(valid_imgs) == 1:
        slide.shapes.add_picture(valid_imgs[0], Inches(img_left), Inches(img_top), height=Inches(img_max_h))
    elif len(valid_imgs) >= 2:
        # Stack 2 images
        h = 2.4
        slide.shapes.add_picture(valid_imgs[0], Inches(img_left), Inches(img_top), width=Inches(img_max_w))
        slide.shapes.add_picture(valid_imgs[1], Inches(img_left), Inches(img_top + h + 0.2), width=Inches(img_max_w))

def draw_flowchart(slide, steps, start_top=2.0):
    # draws horizontal flow
    box_w = 2.2
    box_h = 1.0
    left_offset = 0.5
    spacing = 0.6
    
    for i, step in enumerate(steps):
        l = left_offset + (i * (box_w + spacing))
        # Draw Box
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(start_top), Inches(box_w), Inches(box_h))
        rect.fill.solid()
        rect.fill.fore_color.rgb = NAVY
        rect.line.color.rgb = NAVY
        tf = rect.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = step
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
        # Draw Arrow
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(l + box_w + 0.1), Inches(start_top + 0.35), Inches(spacing - 0.2), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = BLUE
            arrow.line.color.rgb = BLUE

def add_architecture_slide(title, desc_text, flow_steps):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_footer(slide)
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.color.rgb = NAVY
    p.font.bold = True
    
    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12), Inches(1))
    tf = txBox2.text_frame
    p2 = tf.paragraphs[0]
    p2.text = desc_text
    p2.font.size = Pt(20)
    p2.font.color.rgb = DARK_GRAY
    
    draw_flowchart(slide, flow_steps, start_top=3.0)

# --- 3. SLIDE CREATION ---
add_title_slide()

add_content_slide(
    "The Evolution of Event Management",
    [
        "Traditional Event Management: Manual spreadsheets, static ticket systems, and reactive communication.",
        "Our Platform's Approach: A decoupled, microservices-oriented topology.",
        "Agentic Operations: Integrating autonomous AI agents for venue optimization, speaker scheduling, and live incident response.",
        "Full Enterprise Deployment: Ensuring data integrity via Prisma, scalability via Node/FastAPI, and real-time push via Socket.io."
    ]
)

add_architecture_slide(
    "Decoupled Microservices Architecture",
    "The platform separates concerns across Presentation, Orchestration, Intelligence, and Persistence layers.",
    ["React Frontends\n(Student, Admin, Sponsor)", "Node.js / Express\n(Orchestration & Socket.io)", "Python / FastAPI\n(AI Agents & LLM)", "PostgreSQL & Prisma\n(Database Persistence)"]
)

# Milestone 1
add_section_header("Milestone 1", "Intelligent Registration & Attendee Management")

add_content_slide(
    "Intelligent Registration Workflows",
    [
        "AI-Powered Spam & Identity Validation: Zero-shot prompting via Gemini Flash blocks fraudulent registrations.",
        "Dynamic Waitlist Engine: Prevents capacity race conditions using strict database transactions.",
        "Automated Email Ticketing: Dispatches dynamic QR codes upon successful registration.",
        "Real-Time VIP Alerts: Socket.io triggers sub-second toasts to staff when a VIP arrives."
    ]
)

add_dashboard_showcase(
    "Milestone 1: Live Dashboards & Check-in",
    ["Student Portal Registration provides dynamic forms.", "Admin Analytics render live pie charts via Recharts.", "QR Scanner Module triggers real-time VIP alerts."],
    [get_img("M1", -1), get_img("M1", -3)]
)

# Milestone 2
add_section_header("Milestone 2", "Agentic Venue & Speaker Operations")

add_architecture_slide(
    "Agentic Optimization Workflow",
    "Transitioning from simple validation to goal-oriented operational planning.",
    ["Context Retrieval\n(Venues / Speakers)", "Hard Constraint Filtering\n(Capacity, Time)", "LLM Soft Scoring\n(Equipment, Expertise)", "Admin Approval\n& Persistence"]
)

add_dashboard_showcase(
    "Milestone 2: Optimization Dashboards",
    ["Venue Agent: Explains why a venue is the best fit.", "Speaker Agent: Semantically matches bios to session topics.", "Conflict Checking: Blocks overlapping schedules deterministically."],
    [get_img("M2", -1), get_img("M2", -2)]
)

# Milestone 3
add_section_header("Milestone 3", "Live Execution: Sponsorship & Incident Agents")

add_architecture_slide(
    "Live Incident Triage Pipeline",
    "Parsing vague, free-text emergency reports into actionable team alerts.",
    ["Student Free-Text\nIncident Report", "LLM Severity &\nCategory Triage", "Socket.io 'new_alert'\nEmission", "Admin Kanban Board\nLive Update"]
)

add_dashboard_showcase(
    "Milestone 3: Incident & Sponsor Portals",
    ["Incident Management: Real-time Kanban board with AI-assigned urgency.", "Sponsor ROI Tracking: Brand-new portal computing live ROI metrics based on booth traffic.", "Real-Time Adjustments: Zero refresh required for admins."],
    [get_img("M3", -1), get_img("M3", -4)]
)

# Milestone 4
add_section_header("Milestone 4", "Event Intelligence & Enterprise Orchestration")

add_content_slide(
    "Event Intelligence Engine & Orchestration",
    [
        "Proactive Monitoring: Automatically detects anomalies like hall capacity limits.",
        "Actionable Recommendations: Provides one-click 'Accept / Dismiss' solutions to executives.",
        "Multi-Agent Orchestrator: Decomposes complex problems and delegates them to the Venue and Sponsor agents in parallel.",
        "RBAC Hardening: Strict authorization closures to ensure sponsors cannot access admin layers."
    ]
)

add_dashboard_showcase(
    "Milestone 4: Executive Command Center",
    ["Command Center: Side-by-side view of anomalies and one-click AI solutions.", "Agents Fleet Visualization: Inspectable JSON schemas and reasoning flows.", "Full Audit Logging: Records exactly who approved which AI recommendation."],
    [get_img("M4", -2), get_img("M4", -3)]
)

add_content_slide(
    "Conclusion",
    [
        "Delivered a cohesive, 5-service enterprise platform.",
        "Successfully merged generative AI with deterministic software engineering.",
        "Achieved true Agentic operational intelligence, proving that AI can safely manage high-stakes, real-world events."
    ]
)

output_path = r"C:\Users\vinay\Desktop\InfosysProject1\Agentic_AI_Event_Management_Enterprise.pptx"
prs.save(output_path)
print(f"Enterprise PPTX successfully saved to {output_path}")

# Cleanup
for f in glob.glob(os.path.join(output_img_dir, "*")):
    os.remove(f)
os.rmdir(output_img_dir)
print("Temporary images cleaned up.")
